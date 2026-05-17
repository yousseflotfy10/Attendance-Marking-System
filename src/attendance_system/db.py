from __future__ import annotations

import hashlib
from datetime import date, datetime
from pathlib import Path
from typing import Any, List, Optional

try:
    from openpyxl import Workbook, load_workbook
except Exception:
    Workbook = None
    load_workbook = None

from .config import DATABASE_PATH, DEFAULT_ADMIN_PASSWORD, DEFAULT_ADMIN_USERNAME
from .entities import AttendanceRecord, Student


STUDENTS_SHEET = "students"
ATTENDANCE_SHEET = "attendance"
ADMINS_SHEET = "admins"

STUDENTS_HEADERS = ["id", "name", "department", "created_at"]
ATTENDANCE_HEADERS = ["id", "student_id", "attendance_date", "attendance_time", "status", "created_at"]
ADMINS_HEADERS = ["id", "username", "password_hash", "created_at"]


class AttendanceDatabase:
    def __init__(self, database_path: Path = DATABASE_PATH) -> None:
        self.database_path = Path(database_path)
        self.database_path.parent.mkdir(parents=True, exist_ok=True)
        self.initialize()

    @staticmethod
    def _hash_password(password: str) -> str:
        return hashlib.sha256(password.encode("utf-8")).hexdigest()

    @staticmethod
    def _now_iso() -> str:
        return datetime.now().isoformat(timespec="seconds")

    def _require_excel(self) -> None:
        if Workbook is None or load_workbook is None:
            raise RuntimeError("openpyxl is required for Excel storage. Install it with: pip install openpyxl")

    def _create_sheet(self, workbook, title: str, headers: list[str]):
        sheet = workbook.create_sheet(title=title)
        sheet.append(headers)
        return sheet

    def _ensure_workbook(self) -> None:
        self._require_excel()
        if not self.database_path.exists():
            workbook = Workbook()
            if "Sheet" in workbook.sheetnames:
                del workbook["Sheet"]
            self._create_sheet(workbook, STUDENTS_SHEET, STUDENTS_HEADERS)
            self._create_sheet(workbook, ATTENDANCE_SHEET, ATTENDANCE_HEADERS)
            self._create_sheet(workbook, ADMINS_SHEET, ADMINS_HEADERS)
            workbook.save(self.database_path)
            return

        workbook = load_workbook(self.database_path)
        changed = False

        if STUDENTS_SHEET not in workbook.sheetnames:
            self._create_sheet(workbook, STUDENTS_SHEET, STUDENTS_HEADERS)
            changed = True
        if ATTENDANCE_SHEET not in workbook.sheetnames:
            self._create_sheet(workbook, ATTENDANCE_SHEET, ATTENDANCE_HEADERS)
            changed = True
        if ADMINS_SHEET not in workbook.sheetnames:
            self._create_sheet(workbook, ADMINS_SHEET, ADMINS_HEADERS)
            changed = True

        if changed:
            workbook.save(self.database_path)

    def _load(self):
        self._ensure_workbook()
        return load_workbook(self.database_path)

    def _sheet_rows(self, sheet) -> list[dict[str, Any]]:
        headers = [cell.value for cell in sheet[1]]
        rows: list[dict[str, Any]] = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row is None:
                continue
            if all(cell is None for cell in row):
                continue
            rows.append({headers[index]: row[index] for index in range(len(headers))})
        return rows

    @staticmethod
    def _next_id(rows: list[dict[str, Any]]) -> int:
        if not rows:
            return 1
        return int(max(int(row["id"]) for row in rows if row.get("id") is not None) + 1)

    def initialize(self) -> None:
        self._ensure_workbook()
        self._ensure_default_admin()

    def _ensure_default_admin(self) -> None:
        workbook = self._load()
        admins_sheet = workbook[ADMINS_SHEET]
        admins = self._sheet_rows(admins_sheet)

        existing = None
        for row in admins:
            if str(row.get("username", "")).strip() == DEFAULT_ADMIN_USERNAME:
                existing = row
                break

        if existing is None:
            password_hash = self._hash_password(DEFAULT_ADMIN_PASSWORD)
            admins_sheet.append(
                [
                    self._next_id(admins),
                    DEFAULT_ADMIN_USERNAME,
                    password_hash,
                    self._now_iso(),
                ]
            )
            workbook.save(self.database_path)

    def authenticate_admin(self, username: str, password: str) -> bool:
        normalized_username = username.strip()
        if not normalized_username or not password:
            return False

        workbook = self._load()
        admins = self._sheet_rows(workbook[ADMINS_SHEET])
        password_hash = self._hash_password(password)

        for row in admins:
            if str(row.get("username", "")).strip() == normalized_username and str(row.get("password_hash", "")) == password_hash:
                return True
        return False

    def upsert_student(self, name: str, department: str) -> Student:
        normalized_name = name.strip()
        normalized_department = department.strip()
        workbook = self._load()
        students_sheet = workbook[STUDENTS_SHEET]
        students = self._sheet_rows(students_sheet)

        for row_idx, row in enumerate(students, start=2):
            if str(row.get("name", "")).strip().lower() == normalized_name.lower():
                students_sheet.cell(row=row_idx, column=3, value=normalized_department)
                workbook.save(self.database_path)
                return Student(id=int(row["id"]), name=normalized_name, department=normalized_department)

        new_id = self._next_id(students)
        students_sheet.append([new_id, normalized_name, normalized_department, self._now_iso()])
        workbook.save(self.database_path)
        return Student(id=new_id, name=normalized_name, department=normalized_department)

    def get_student_by_name(self, name: str) -> Optional[Student]:
        normalized_name = name.strip().lower()
        workbook = self._load()
        students = self._sheet_rows(workbook[STUDENTS_SHEET])

        for row in students:
            if str(row.get("name", "")).strip().lower() == normalized_name:
                return Student(
                    id=int(row["id"]),
                    name=str(row["name"]),
                    department=str(row["department"]),
                )
        return None

    def list_students(self) -> List[Student]:
        workbook = self._load()
        students = self._sheet_rows(workbook[STUDENTS_SHEET])
        students_sorted = sorted(students, key=lambda row: str(row.get("name", "")).lower())
        return [
            Student(
                id=int(row["id"]),
                name=str(row["name"]),
                department=str(row["department"]),
            )
            for row in students_sorted
        ]

    def mark_attendance(self, student_id: int, status: str = "present") -> bool:
        today = date.today().isoformat()
        now_time = datetime.now().strftime("%H:%M:%S")

        workbook = self._load()
        attendance_sheet = workbook[ATTENDANCE_SHEET]
        attendance_rows = self._sheet_rows(attendance_sheet)

        for row in attendance_rows:
            if int(row.get("student_id", 0)) == int(student_id) and str(row.get("attendance_date", "")) == today:
                return False

        new_id = self._next_id(attendance_rows)
        attendance_sheet.append([new_id, int(student_id), today, now_time, status, self._now_iso()])
        workbook.save(self.database_path)
        return True

    def has_attendance_for_today(self, student_id: int) -> bool:
        today = date.today().isoformat()
        workbook = self._load()
        attendance_rows = self._sheet_rows(workbook[ATTENDANCE_SHEET])
        for row in attendance_rows:
            if int(row.get("student_id", 0)) == int(student_id) and str(row.get("attendance_date", "")) == today:
                return True
        return False

    def get_marked_student_names(self) -> set[str]:
        today = date.today().isoformat()
        workbook = self._load()
        students = self._sheet_rows(workbook[STUDENTS_SHEET])
        attendance_rows = self._sheet_rows(workbook[ATTENDANCE_SHEET])

        id_to_name = {int(row["id"]): str(row["name"]) for row in students}
        marked = set()
        for row in attendance_rows:
            if str(row.get("attendance_date", "")) == today:
                student_name = id_to_name.get(int(row.get("student_id", 0)))
                if student_name:
                    marked.add(student_name)
        return marked

    def list_today_attendance(self) -> List[AttendanceRecord]:
        today = date.today().isoformat()
        workbook = self._load()
        students = self._sheet_rows(workbook[STUDENTS_SHEET])
        attendance_rows = self._sheet_rows(workbook[ATTENDANCE_SHEET])

        students_by_id = {int(row["id"]): row for row in students}
        today_rows = [row for row in attendance_rows if str(row.get("attendance_date", "")) == today]
        today_rows.sort(key=lambda row: str(row.get("attendance_time", "")), reverse=True)

        records: List[AttendanceRecord] = []
        for row in today_rows:
            student_row = students_by_id.get(int(row.get("student_id", 0)))
            if student_row is None:
                continue
            records.append(
                AttendanceRecord(
                    id=int(row["id"]),
                    student_id=int(row["student_id"]),
                    student_name=str(student_row["name"]),
                    department=str(student_row["department"]),
                    attendance_date=str(row["attendance_date"]),
                    attendance_time=str(row["attendance_time"]),
                    status=str(row["status"]),
                )
            )
        return records
